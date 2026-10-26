import sys
from docx import Document

# Read the Word document
doc = Document(r'C:\College\Project\5G\5G_report_midsem.docx')

# Extract all text with UTF-8 encoding
with open(r'C:\College\Project\5G\extracted_report.txt', 'w', encoding='utf-8') as f:
    for para in doc.paragraphs:
        if para.text.strip():
            f.write(para.text + '\n')
    
    # Also extract tables
    for table in doc.tables:
        f.write('\n--- TABLE ---\n')
        for row in table.rows:
            row_text = '\t'.join([cell.text for cell in row.cells])
            f.write(row_text + '\n')

print("Content extracted to extracted_report.txt")
